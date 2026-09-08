"""``extract_table`` — any table as a dataset; the shape is dictated by ``inputs.expected_output``.

The Excel Extractor used to fix only two hardcoded shapes (``facts`` → dates, ``new_wells`` → parameters).
These tests pin the flexible path: the orchestrator names the dataset and its fields, the LLM maps
columns, extraction stays deterministic, rows travel as ``data[<name>]`` + a JSON artifact, and the
result says which requested fields the workbook could not provide.
"""

from __future__ import annotations

import json
from datetime import datetime
from pathlib import Path

import pytest
from openpyxl import Workbook

from mas_agent_kit import ActivityClient, is_dataset, result_text_problems
from mas_agent_kit.dataset import INLINE_ROWS_BYTES

from app import agent as agent_module
from app import agent_tools, dataset_extract as ds
from app.agent import agent
from app.sessions import load_state
from app.tools import execute_tool

EVENTS = [
    # Скважина (number cell → text "101"), Дата (datetime), Мероприятие, Значение (mixed number formats), Комментарий
    (101, datetime(2026, 3, 1), "ГРП", 12.5, "первый этап"),
    (102, datetime(2026, 3, 15), "Перевод в ППД", None, ""),
    ("N-7", datetime(2026, 4, 1), "Смена режима", "1 200,5", "ORAT"),
    (103, "05.04.2026", "ГРП", 8, None),
    (101, datetime(2026, 5, 1), "Остановка", 0, "дубликат скважины — другая дата"),
]
MONTHS = ["01.2026", "02.2026", "03.2026"]
RATES = [("1601", 120.5, 118, 110), ("1735", None, 80, 79.5), ("2012", 45, 44, None)]


def _events_workbook(path: Path) -> None:
    wb = Workbook()
    noise = wb.active
    noise.title = "Примечания"
    noise.append(["Файл подготовлен отделом разработки"])
    noise.append(["Версия", 3])
    ws = wb.create_sheet("Мероприятия")
    ws.append(["Мероприятия по скважинам на 2026 год"])  # merged-style title row
    ws.append([])
    ws.append(["Скважина", "Дата", "Мероприятие", "Значение", "Комментарий"])
    ws.append(["", "", "", "т/сут", ""])  # units row
    for row in EVENTS:
        ws.append(list(row))
    ws.append(["Итого", None, None, 20.5, None])
    lookup = wb.create_sheet("Справочник")
    lookup.append(["Скважина", "Куст"])
    lookup.append(["101", "1"])
    lookup.append(["102", "1"])
    wb.save(path)


def _wide_workbook(path: Path) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Дебиты"
    ws.append(["Скважина", "Группа", *MONTHS])
    for well, *values in RATES:
        ws.append([well, "G1" if well != "2012" else "G2", *values])
    wb.save(path)


def _task(path: Path, expected: dict | None = None, **inputs) -> dict:
    return {
        "case_id": "CASE-T",
        "task_id": "TASK-T",
        "agent_id": "excel_extractor",
        "objective": "Извлеки таблицу мероприятий",
        "handoff_message": "Нужны мероприятия по скважинам: скважина, дата, вид, значение.",
        "inputs": {"excel_path": str(path), **({"expected_output": expected} if expected else {}), **inputs},
        "context": {},
    }


EXPECTED = {
    "datasets": [
        {
            "name": "well_events",
            "description": "мероприятия по скважинам",
            "fields": [{"name": "well", "description": "скважина"}, {"name": "date", "type": "date"}, {"name": "kind"}, {"name": "value", "type": "number"}, {"name": "crew"}],
        }
    ],
    "consumers": [{"agent_id": "schedule_builder", "title": "Schedule Builder", "needs": {"facts": "даты ввода"}}],
}


@pytest.fixture(autouse=True)
def _quiet_activity(monkeypatch, tmp_path):
    monkeypatch.setenv("SESSION_DIR", str(tmp_path))
    seen: list[tuple[str, str]] = []
    monkeypatch.setattr(ActivityClient, "event", lambda self, kind, message, **kw: seen.append((kind, message)) or True)
    agent_module._SEEN_EVENTS = seen  # type: ignore[attr-defined]
    yield seen


def _open(tmp_path: Path, expected: dict | None = EXPECTED, **inputs):
    path = tmp_path / "events.xlsx"
    _events_workbook(path)
    opened = agent.open_session(_task(path, expected, **inputs))
    assert opened["ok"] is True, opened
    table = next(t for t in opened["inspect"]["tables"] if t["sheet"] == "Мероприятия")
    return opened, table


def test_open_session_hands_the_expected_output_to_the_model(tmp_path: Path) -> None:
    opened, table = _open(tmp_path)
    assert opened["expected_output"]["datasets"][0]["name"] == "well_events"
    assert opened["expected_output"]["consumers"][0]["needs"] == {"facts": "даты ввода"}
    assert table["columns"] == ["Скважина", "Дата", "Мероприятие", "Значение", "Комментарий"], "units row and title are not columns"
    opened_plain, _ = _open(tmp_path, expected=None)
    assert opened_plain["expected_output"] == {}


def test_extract_table_maps_fields_types_rows_and_reports_missing_requested_fields(tmp_path: Path) -> None:
    opened, table = _open(tmp_path)
    state = load_state(opened["session_id"])
    done = execute_tool(
        state,
        "extract_table",
        {"table_id": table["table_id"], "name": "well_events", "columns": {"well": "Скважина", "date": "Дата", "kind": "Мероприятие", "value": "Значение"}, "title": "Мероприятия по скважинам"},
    )
    assert done["ok"] is True, done
    assert done["status"] == "completed" and done["dataset"] == "well_events" and done["row_count"] == 5
    assert done["issues"]["requested_fields_missing"] == ["crew"], "the orchestrator asked for a field the table does not have"
    assert done["next_step"].startswith("Результат зафиксирован"), "well_events was the only expected dataset"

    result = agent.result(load_state(opened["session_id"]))
    assert result["status"] == "completed"
    entry = result["data"]["well_events"]
    assert is_dataset(entry) and entry["row_count"] == 5 and entry["title"] == "Мероприятия по скважинам"
    assert [(f["name"], f["type"], f["source_column"]) for f in entry["fields"]] == [
        ("well", "text", "Скважина"),
        ("date", "date", "Дата"),
        ("kind", "text", "Мероприятие"),
        ("value", "number", "Значение"),
    ]
    rows = entry["rows"]
    assert rows[0] == {"well": "101", "date": "2026-03-01", "kind": "ГРП", "value": 12.5}, "101.0 → '101', datetime → ISO"
    assert rows[1] == {"well": "102", "date": "2026-03-15", "kind": "Перевод в ППД"}, "empty cells are omitted, not null"
    assert rows[2]["value"] == 1200.5, "'1 200,5' is a number"
    assert rows[3]["date"] == "2026-04-05", "dd.mm.yyyy text is a date"
    assert rows[4]["well"] == "101" and rows[4]["value"] == 0, "zero is a value; duplicates stay without key_field"
    assert "Итого" not in json.dumps(rows, ensure_ascii=False), "the total row is not data"
    assert entry["source"]["sheet"] == "Мероприятия" and entry["source"]["table_id"] == table["table_id"]
    assert entry["issues"] == {"requested_fields_missing": ["crew"]}
    assert result["data"]["expected"] == {"requested": ["well_events"], "produced": ["well_events"], "missing": []}
    assert result["message"] == "Извлёк набор данных «Мероприятия по скважинам»: 5 строк, 4 поля. В таблице не нашлось 1 запрошенного поля."
    assert result_text_problems(result) == []
    assert result["artifacts"] == {"excel_session": opened["session_id"]}, "no Activity → no artifact card, rows inline"
    assert any(msg.startswith("Извлёк набор данных") for _, msg in agent_module._SEEN_EVENTS)  # type: ignore[attr-defined]


def test_extract_table_uploads_the_rows_as_an_artifact_and_returns_its_card(tmp_path: Path, monkeypatch) -> None:
    uploads: list[dict] = []

    def fake_upload(self, artifact_id, filename, content, mime_type="application/octet-stream", *, summary=""):
        uploads.append({"artifact_id": artifact_id, "filename": filename, "body": json.loads(content.decode("utf-8")), "mime": mime_type, "summary": summary})
        return {"artifact_id": artifact_id + "_1", "filename": filename, "role": "attachment", "kind": "deliverable", "producer": "excel_extractor", "bytes": len(content), "download_path": "/x"}

    monkeypatch.setattr(ActivityClient, "upload", fake_upload)
    opened, table = _open(tmp_path, activity_base_url="http://activity.test")
    state = load_state(opened["session_id"])
    done = execute_tool(state, "extract_table", {"table_id": table["table_id"], "name": "well_events", "columns": {"well": "Скважина", "date": "Дата"}})
    assert done["ok"] is True, done
    assert uploads and uploads[0]["artifact_id"] == "dataset_well_events" and uploads[0]["filename"] == "well_events.json" and uploads[0]["mime"] == "application/json"
    assert uploads[0]["body"]["rows"][0] == {"well": "101", "date": "2026-03-01"} and uploads[0]["body"]["name"] == "well_events"
    result = agent.result(load_state(opened["session_id"]))
    card = result["artifacts"]["dataset_well_events_1"]
    assert card["producer"] == "excel_extractor" and "download_path" not in card
    assert result["data"]["well_events"]["artifact_id"] == "dataset_well_events_1", "the id Activity assigned, not the requested one"
    assert result["data"]["well_events"]["rows"], "small tables stay inline as well"


def test_large_dataset_keeps_a_preview_inline(tmp_path: Path) -> None:
    path = tmp_path / "big.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.title = "История"
    ws.append(["Скважина", "Дата", "Комментарий"])
    for i in range(600):
        ws.append([1000 + i, datetime(2026, 1, 1), "довольно длинный комментарий к строке номер %d" % i])
    wb.save(path)
    opened = agent.open_session(_task(path, None))
    state = load_state(opened["session_id"])
    table_id = opened["inspect"]["tables"][0]["table_id"]
    done = execute_tool(state, "extract_table", {"table_id": table_id, "name": "history"})
    assert done["ok"] is True and done["row_count"] == 600
    entry = agent.result(load_state(opened["session_id"]))["data"]["history"]
    assert "rows" not in entry and len(entry["preview"]) == 5
    assert len(json.dumps(entry, ensure_ascii=False)) < INLINE_ROWS_BYTES
    assert [f["name"] for f in entry["fields"]] == ["skvazhina", "data", "kommentarij"], "no columns given → every column, field names from headers"


def test_wide_table_is_unpivoted_into_long_rows(tmp_path: Path) -> None:
    path = tmp_path / "rates.xlsx"
    _wide_workbook(path)
    opened = agent.open_session(_task(path, {"datasets": [{"name": "oil_rates", "fields": [{"name": "well", "type": "text"}, {"name": "date", "type": "date"}, {"name": "rate", "type": "number"}]}]}))
    state = load_state(opened["session_id"])
    table = opened["inspect"]["tables"][0]
    assert table["columns"] == ["Скважина", "Группа", *MONTHS]
    # «rest» would also melt the group column — the model names the wide columns explicitly.
    done = execute_tool(
        state,
        "extract_table",
        {"table_id": table["table_id"], "name": "oil_rates", "columns": {"well": "Скважина", "group": "Группа"}, "unpivot": {"columns": MONTHS, "name_field": "date", "value_field": "rate"}},
    )
    assert done["ok"] is True, done
    assert done["row_count"] == 7, "3 wells × 3 months minus 2 empty cells"
    entry = agent.result(load_state(opened["session_id"]))["data"]["oil_rates"]
    assert [(f["name"], f["type"]) for f in entry["fields"]] == [("well", "text"), ("group", "text"), ("date", "date"), ("rate", "number")]
    assert entry["fields"][2]["source_column"].startswith("заголовки колонок: 01.2026")
    assert entry["rows"][0] == {"well": "1601", "group": "G1", "date": "2026-01-01", "rate": 120.5}
    assert entry["rows"][3] == {"well": "1735", "group": "G1", "date": "2026-02-01", "rate": 80}
    assert not entry.get("issues"), "no requested field is missing: well, date, rate are all there"


def test_filters_key_dedupe_and_declared_types(tmp_path: Path) -> None:
    opened, table = _open(tmp_path, expected=None)
    state = load_state(opened["session_id"])
    done = execute_tool(
        state,
        "extract_table",
        {
            "table_id": table["table_id"],
            "name": "frac_jobs",
            "columns": {"well": "Скважина", "date": "Дата", "value": "Значение"},
            "filters": [{"field": "Мероприятие", "operator": "eq", "value": "ГРП"}],
            "key_field": "well",
            "types": {"value": "text"},
        },
    )
    assert done["ok"] is True, done
    entry = agent.result(load_state(opened["session_id"]))["data"]["frac_jobs"]
    assert entry["rows"] == [{"well": "101", "date": "2026-03-01", "value": "12.5"}, {"well": "103", "date": "2026-04-05", "value": "8"}]
    assert entry["issues"] == {"filtered_out": 3}
    assert entry["fields"][2]["type"] == "text", "a declared type wins over inference"

    # Dedupe by key: the second «101» row is dropped and counted.
    done = execute_tool(state, "extract_table", {"table_id": table["table_id"], "name": "last_events", "columns": {"well": "Скважина", "kind": "Мероприятие"}, "key_field": "well"})
    entry = agent.result(load_state(opened["session_id"]))["data"]["last_events"]
    assert [r["well"] for r in entry["rows"]] == ["101", "102", "N-7", "103"] and entry["issues"] == {"duplicates": 1}


def test_extract_table_argument_errors_are_for_the_model(tmp_path: Path) -> None:
    opened, table = _open(tmp_path)
    state = load_state(opened["session_id"])
    reserved = execute_tool(state, "extract_table", {"table_id": table["table_id"], "name": "facts", "columns": {"well": "Скважина", "date": "Дата"}})
    assert reserved["ok"] is False and reserved["code"] == "name_reserved" and reserved["use_instead"] == "extract_commissioning"
    bad_name = execute_tool(state, "extract_table", {"table_id": table["table_id"], "name": "мероприятия"})
    assert bad_name["ok"] is False and bad_name["code"] == "spec_incomplete" and bad_name["missing"] == ["name"]
    wrong_col = execute_tool(state, "extract_table", {"table_id": table["table_id"], "name": "x_events", "columns": {"well": "Скв", "date": "Нет такой"}})
    assert wrong_col["ok"] is False and wrong_col["code"] == "column_not_found"
    assert wrong_col["not_found"] == {"date": "Нет такой"} and wrong_col["available_columns"] == table["columns"], "unique prefix «Скв» resolves, «Нет такой» does not"
    opened2, table2 = _open(tmp_path)
    state2 = load_state(opened2["session_id"])
    empty = execute_tool(state2, "extract_table", {"table_id": table2["table_id"], "name": "none_events", "columns": ["Скважина"], "filters": [{"field": "Мероприятие", "operator": "eq", "value": "Нет"}]})
    assert empty["ok"] is False and empty["code"] == "no_rows" and empty["filtered_out"] == 5
    no_table = execute_tool(state2, "extract_table", {"table_id": "tbl_nope", "name": "y_events"})
    assert no_table["ok"] is False and no_table["code"] == "table_not_found"
    assert agent.result(load_state(opened["session_id"]))["status"] == "needs_input", "nothing was fixed by failed calls"
    for envelope in (reserved, bad_name, wrong_col, empty, no_table):
        assert "error" not in envelope, "n8n retries a tool item whose json has an error key (A15)"


def test_normalize_args_accepts_n8n_and_function_calling_vocabulary() -> None:
    normalized = agent.normalize_args(
        "extract_table",
        {"dataset": "well_events", "mapping": {"well": "Скважина"}, "filters": {"0": {"field": "Мероприятие", "operator": "eq", "value": "ГРП"}}, "key": "well", "unpivot": '{"columns": "rest"}'},
    )
    assert normalized["name"] == "well_events" and normalized["columns"] == {"well": "Скважина"} and normalized["key_field"] == "well"
    assert normalized["filters"] == [{"field": "Мероприятие", "operator": "eq", "value": "ГРП"}]
    assert normalized["unpivot"] == '{"columns": "rest"}', "JSON text is parsed by the tool, not here"
    by_column = agent.normalize_args("extract_table", {"filters": {"Мероприятие": "ГРП"}})
    assert by_column["filters"] == [{"field": "Мероприятие", "operator": "eq", "value": "ГРП"}]


def test_expected_datasets_drive_next_step_and_the_final_message(tmp_path: Path) -> None:
    """Two expected datasets: after the first the model is told what is still missing; a result that ends
    without the second says so in prose and in ``issues`` (the orchestrator decides what to do)."""
    expected = {"datasets": [{"name": "well_events", "fields": [{"name": "well"}]}, {"name": "well_notes", "description": "примечания по скважинам", "fields": [{"name": "well"}, {"name": "note"}]}]}
    opened, table = _open(tmp_path, expected=expected)
    state = load_state(opened["session_id"])
    first = execute_tool(state, "extract_table", {"table_id": table["table_id"], "name": "well_events", "columns": {"well": "Скважина", "kind": "Мероприятие"}})
    assert first["ok"] is True
    assert "well_notes (примечания по скважинам)" in first["next_step"] and "extract_table" in first["next_step"]
    partial = agent.result(load_state(opened["session_id"]))
    assert partial["status"] == "completed"
    assert partial["data"]["expected"] == {"requested": ["well_events", "well_notes"], "produced": ["well_events"], "missing": ["well_notes"]}
    assert partial["issues"] == [{"type": "expected_output_missing", "datasets": ["well_notes"]}]
    assert partial["message"].endswith("Один из запрошенных наборов данных не извлечён.")
    assert result_text_problems(partial) == []
    second = execute_tool(state, "extract_table", {"table_id": table["table_id"], "name": "well_notes", "columns": {"well": "Скважина", "note": "Комментарий"}})
    assert second["ok"] is True and second["next_step"].startswith("Результат зафиксирован")
    full = agent.result(load_state(opened["session_id"]))
    assert full["issues"] == [] and "не извлечён" not in full["message"]
    assert full["data"]["expected"]["missing"] == [] and set(full["data"]) >= {"well_events", "well_notes"}
    assert full["message"].startswith("Извлёк набор данных: 5 строк, 2 поля. Извлёк набор данных «примечания по скважинам»")


def test_commissioning_and_table_parts_accumulate_and_facts_count_as_produced(tmp_path: Path) -> None:
    expected = {"datasets": [{"name": "facts", "fields": [{"name": "well"}, {"name": "date"}]}, {"name": "well_events", "fields": [{"name": "well"}]}]}
    opened, table = _open(tmp_path, expected=expected)
    state = load_state(opened["session_id"])
    dates = execute_tool(state, "extract_commissioning", {"table_id": table["table_id"], "well_column": "Скважина", "date_column": "Дата"})
    assert dates["ok"] is True and "well_events" in dates["next_step"], "facts fixed, well_events still expected"
    events = execute_tool(state, "extract_table", {"table_id": table["table_id"], "name": "well_events", "columns": {"well": "Скважина", "kind": "Мероприятие"}})
    assert events["ok"] is True
    result = agent.result(load_state(opened["session_id"]))
    assert result["data"]["facts"][0] == {"well": "101", "date": "2026-03-01"} and is_dataset(result["data"]["well_events"])
    assert result["data"]["expected"]["missing"] == [] and result["issues"] == []
    assert result["message"].startswith("Даты ввода: 4 скважины")


def test_value_normalisation_rules() -> None:
    assert ds.as_date("01.2026") == ("2026-01-01", True) and ds.as_date(3.1416) == ("3.1416", False), "month.year only for header text"
    assert ds.as_date(46023, allow_serial=True) == ("2026-01-01", True) and ds.as_date(46023) == ("46023", False), "serials only for declared dates"
    assert ds.as_number("1 200,5") == (1200.5, True) and ds.as_number("12") == (12, True) and ds.as_number("abc") == ("abc", False) and ds.as_number(True) == (None, False)
    assert ds.as_text(101.0) == "101" and ds.as_text("N-7 ") == "N-7" and ds.as_text(True) == "true"
    assert ds.infer_type([45000, 46000, 47000]) == "number", "a column of numbers is never a column of dates"
    assert ds.infer_type(["2026-01-01T00:00:00", "13.02.2026", "нет данных"]) == "text", "2 of 3 (67 %) is below the 80 % agreement"
    assert ds.infer_type(["2026-01-01T00:00:00", "13.02.2026", "14.02.2026", "15.02.2026", "x"]) == "date"
    assert ds.infer_type([True, False]) == "boolean" and ds.infer_type([]) == "text"
    used: set[str] = set()
    assert [ds.field_name_from_header(h, used) for h in ("Дебит нефти, т/сут", "Дебит нефти", "2026", "")] == ["debit_nefti_t_sut", "debit_nefti", "col_2026", "col"]
