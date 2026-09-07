"""Excel Extractor agent: session over the case's Excel attachments → LLM picks table/columns → deterministic extraction.

Contract (n8n «Agent — Excel Extractor» workflow):

* ``open_session(agent_task)`` loads every Excel attachment of the case into one session and returns a
  compact **inventory** (files, sheets, tables with columns and two sample rows). No decisions are
  taken here — no keyword heuristics decide which table or column holds the wells or the dates.
* The LLM reads the inventory and calls ``extract_commissioning(table_id, well_column, date_column)``
  or ``extract_well_parameters(table_id, well_column, mapping)``; both are ordinary registry tools
  (``/agent-tools/{name}``) that validate the choice, extract rows deterministically and store the
  agent result in the session. ``ask_engineer`` is the only way to ask a human (prose, validated).
* ``session_result(session_id)`` returns the stored result (or a human-readable needs_input when the
  LLM finished without extracting anything).

No SCHEDULE text is ever written here.
"""

from __future__ import annotations

import hashlib
import json
import logging
import os
import re
from datetime import date, datetime
from pathlib import Path
from typing import Any
from urllib.request import Request, urlopen

from .excel_tools import _record, _rows_for_table, _schema, _table
from .sessions import init_state, load_state, locked_session, new_session_id, save_state, session_dir, session_file
from .tools import ToolError, execute_tool, tool
from . import excel_tools as _excel_tools  # noqa: F401  # register detect_tables / query_table
from . import state_tools as _state_tools  # noqa: F401

logger = logging.getLogger(__name__)
ACTIVITY = os.getenv("ACTIVITY_BASE_URL", "").rstrip("/")
EXCEL_SUFFIXES = (".xlsx", ".xlsm", ".xltx", ".xltm", ".xls")
MERGEABLE_SUFFIXES = (".xlsx", ".xlsm", ".xltx", ".xltm")
MAX_TOOL_REPEATS = 3
# Canonical engineering facts for a new well (Schedule Builder renders WELSPECS / COMPDATMD / WCONPROD
# from them; see schedule-builder-service timeline_ops.NEW_WELL_FACT_COLUMNS). Anything else in the
# row is passed through under its original header.
NEW_WELL_FIELDS = (
    "well",
    "date",
    "group",
    "phase",
    "i",
    "j",
    "md_top",
    "md_bot",
    "diameter",
    "control",
    "rate",
    "bhp",
    "thp",
    "vfp_table",
    "welltrack_include",
)


# --------------------------------------------------------------------------------------------------
# Activity events
# --------------------------------------------------------------------------------------------------


def _post_event(case_id: str, payload: dict[str, Any]) -> None:
    base = str(payload.pop("_activity_base", "") or ACTIVITY).rstrip("/")
    if not case_id or not base:
        return
    try:
        req = Request(
            f"{base}/cases/{case_id}/events",
            data=json.dumps(payload, ensure_ascii=False).encode("utf-8"),
            headers={"Content-Type": "application/json"},
            method="POST",
        )
        urlopen(req, timeout=8).read()
    except Exception:
        logger.exception("failed to emit agent event case_id=%s kind=%s", case_id, payload.get("kind"))


def _emit(
    case_id: str,
    *,
    kind: str,
    task_id: str,
    activity: str,
    status_message: str,
    status: str = "",
    payload: dict[str, Any] | None = None,
) -> None:
    body: dict[str, Any] = {
        "kind": kind,
        "actor": "excel_extractor",
        "agent_id": "excel_extractor",
        "task_id": task_id,
        "status_message": status_message,
        "_activity_base": activity,
    }
    if status:
        body["status"] = status
    if payload:
        body["payload"] = payload
    _post_event(case_id, body)


def _emit_from_state(state: dict[str, Any], message: str, *, kind: str = "agent.progress", status: str = "") -> None:
    payload = state.get("payload") if isinstance(state.get("payload"), dict) else {}
    case_id = str(payload.get("case_id") or "")
    if not case_id:
        return
    _emit(
        case_id,
        kind=kind,
        task_id=str(payload.get("task_id") or ""),
        activity=str(payload.get("activity_base") or ACTIVITY),
        status_message=message,
        status=status,
    )


def emit_tool_progress(state: dict[str, Any], tool_name: str, tool_result: dict[str, Any]) -> None:
    """Live Activity lines when the n8n AI Agent calls excel-tools HTTP tools."""
    inner = tool_result.get("result") if isinstance(tool_result.get("result"), dict) else {}
    message = ""
    if tool_name == "workbook_introspect":
        sheets = inner.get("sheets") if isinstance(inner.get("sheets"), list) else []
        message = f"Смотрю структуру Excel: листов {len(sheets)}"
    elif tool_name == "detect_tables":
        tables = inner.get("tables") or inner.get("detected_tables") or []
        n = len(tables) if isinstance(tables, list) else 0
        message = f"Нашёл таблиц: {n}"
    elif tool_name == "match_tables":
        message = "Подбираю таблицу под запрос"
    elif tool_name == "describe_table":
        message = "Смотрю колонки таблицы"
    elif tool_name == "list_column_values":
        message = "Читаю значения колонки"
    elif tool_name == "query_table":
        rows = inner.get("preview") or inner.get("rows") or inner.get("records") or []
        n = inner.get("row_count") if isinstance(inner.get("row_count"), int) else (len(rows) if isinstance(rows, list) else 0)
        message = f"Читаю таблицу: {n} строк"
    elif tool_name == "sheet_preview":
        message = "Смотрю фрагмент листа"
    elif tool_name in {"extract_commissioning", "extract_well_parameters", "ask_engineer"}:
        # These tools emit their own precise progress line (counts, wells) or a question.
        if not tool_result.get("ok", True):
            err = tool_result.get("error") if isinstance(tool_result.get("error"), dict) else {}
            code = str(err.get("code") or tool_result.get("error") or "")
            message = {
                "column_not_found": "Уточняю колонки таблицы",
                "table_not_found": "Уточняю таблицу",
                "column_not_dates": "Выбранная колонка не похожа на даты — подбираю другую",
                "spec_incomplete": "Уточняю, какую таблицу и колонки взять",
                "question_not_human": "Формулирую вопрос инженеру",
            }.get(code, "")
    if not message:
        return
    _emit_from_state(state, message)


# --------------------------------------------------------------------------------------------------
# Excel attachments → one session
# --------------------------------------------------------------------------------------------------


def _fetch_artifact(url: str) -> tuple[bytes, str]:
    """Download an Activity artifact; return (bytes, filename from Content-Disposition or '')."""
    with urlopen(url, timeout=30) as resp:
        data = resp.read()
        disposition = str(resp.headers.get("Content-Disposition") or "")
    name = ""
    m = re.search(r"filename\*=UTF-8''([^;]+)", disposition)
    if m:
        from urllib.parse import unquote

        name = unquote(m.group(1).strip())
    else:
        m = re.search(r'filename="?([^";]+)"?', disposition)
        if m:
            name = m.group(1).strip()
    return data, name


def _fetch_case_state(activity: str, case_id: str) -> dict[str, Any]:
    if not activity or not case_id:
        return {}
    try:
        with urlopen(f"{activity}/cases/{case_id}/state", timeout=15) as resp:
            packet = json.loads(resp.read().decode("utf-8", errors="replace"))
    except Exception:
        return {}
    state = packet.get("state") if isinstance(packet, dict) else None
    return state if isinstance(state, dict) else {}


def _task_activity(task: dict[str, Any]) -> str:
    inputs = task.get("inputs") if isinstance(task.get("inputs"), dict) else {}
    return str(inputs.get("activity_base_url") or ACTIVITY).rstrip("/")


def _is_excel_name(name: Any) -> bool:
    return str(name or "").lower().endswith(EXCEL_SUFFIXES)


def _walk_excel_cards(node: Any, out: dict[str, dict[str, Any]], key: str = "") -> None:
    """Collect Excel artifact cards from a nested or flat artifacts view (excel, excel_1, attachments…)."""
    if isinstance(node, list):
        for item in node:
            _walk_excel_cards(item, out)
        return
    if not isinstance(node, dict):
        return
    is_card = "artifact_id" in node or isinstance(node.get("filename"), str)
    if is_card:
        aid = str(node.get("artifact_id") or key or "").strip()
        role = str(node.get("role") or "")
        is_excel = role == "excel" or (
            not role and (bool(re.fullmatch(r"excel(_\d+)?", aid)) or _is_excel_name(node.get("filename")))
        )
        if is_excel and aid and aid not in out:
            out[aid] = {"artifact_id": aid, "filename": str(node.get("filename") or "")}
        return
    for child_key, value in node.items():
        _walk_excel_cards(value, out, str(child_key))


def _excel_cards(task: dict[str, Any]) -> list[dict[str, Any]]:
    """Every Excel workbook attached to the case: explicit path, explicit URL, artifact cards, case state."""
    inputs = task.get("inputs") if isinstance(task.get("inputs"), dict) else {}
    case_id = str(task.get("case_id") or "")
    activity = _task_activity(task)
    cards: list[dict[str, Any]] = []
    path = inputs.get("excel_path")
    if path and Path(str(path)).is_file():
        cards.append({"artifact_id": "excel", "filename": Path(str(path)).name, "path": str(path)})
    url = inputs.get("excel_url") or inputs.get("artifact_url")
    if url:
        cards.append({"artifact_id": "excel", "filename": str(inputs.get("excel_filename") or ""), "url": str(url)})
    if cards:
        return cards
    found: dict[str, dict[str, Any]] = {}
    _walk_excel_cards(inputs.get("artifacts"), found)
    if activity and case_id:
        _walk_excel_cards(_fetch_case_state(activity, case_id).get("artifacts"), found)
    ids = inputs.get("artifact_ids") if isinstance(inputs.get("artifact_ids"), list) else []
    for aid in ids:
        aid = str(aid or "").strip()
        if re.fullmatch(r"excel(_\d+)?", aid) and aid not in found:
            found[aid] = {"artifact_id": aid, "filename": ""}
    if activity and case_id and "excel" not in found:
        # The first workbook of a case always has the fixed id `excel`; older orchestrator states may
        # list only later workbooks (or none at all). Probe it — a missing artifact is skipped on download.
        found["excel"] = {"artifact_id": "excel", "filename": ""}

    def order(card: dict[str, Any]) -> tuple[int, str]:
        aid = card["artifact_id"]
        m = re.fullmatch(r"excel_(\d+)", aid)
        return (int(m.group(1)) if m else 0, aid)

    for card in sorted(found.values(), key=order):
        if activity and case_id:
            card["url"] = f"{activity}/cases/{case_id}/artifacts/{card['artifact_id']}"
        cards.append(card)
    return cards


def _download_cards(cards: list[dict[str, Any]]) -> list[tuple[str, bytes]]:
    files: list[tuple[str, bytes]] = []
    errors: list[str] = []
    for card in cards:
        data: bytes | None = None
        name = str(card.get("filename") or "")
        if card.get("path"):
            data = Path(str(card["path"])).read_bytes()
            name = name or Path(str(card["path"])).name
        elif card.get("url"):
            try:
                data, header_name = _fetch_artifact(str(card["url"]))
            except Exception as exc:  # noqa: BLE001
                logger.warning("excel artifact %s skipped: %s", card.get("artifact_id"), exc)
                errors.append(f"{card.get('artifact_id')}: {exc}")
                continue
            name = name or header_name
        if not data:
            continue
        if not _is_excel_name(name):
            name = (name or str(card.get("artifact_id") or "workbook")) + ".xlsx"
        files.append((name, data))
    if not files and errors:
        raise FileNotFoundError("excel artifact fetch failed: " + "; ".join(errors)[:400])
    return files


def _safe_sheet_title(title: str, used: set[str]) -> str:
    clean = re.sub(r"[\[\]:*?/\\]", " ", title).strip() or "Sheet"
    base = clean[:31]
    candidate = base
    n = 2
    while candidate in used:
        suffix = f" ({n})"
        candidate = base[: 31 - len(suffix)] + suffix
        n += 1
    used.add(candidate)
    return candidate


def _merge_workbooks(files: list[tuple[str, bytes]]) -> tuple[str, bytes, list[dict[str, Any]]]:
    """One session file for all attachments.

    A single workbook is stored as-is. Several .xlsx are merged into one workbook (values only) with
    sheet titles ``<sheet> (<file stem>)`` so the inventory tells the LLM which file a table came from.
    """
    if len(files) == 1:
        name, data = files[0]
        return name, data, [{"file": name, "sheets": [], "merged": False}]
    import io

    from openpyxl import Workbook, load_workbook

    target = Workbook()
    target.remove(target.active)
    used: set[str] = set()
    sources: list[dict[str, Any]] = []
    for name, data in files:
        if not name.lower().endswith(MERGEABLE_SUFFIXES):
            sources.append({"file": name, "sheets": [], "merged": False, "skipped": "only .xlsx workbooks are merged"})
            continue
        try:
            src = load_workbook(io.BytesIO(data), data_only=True)
        except Exception as exc:  # noqa: BLE001
            sources.append({"file": name, "sheets": [], "merged": False, "skipped": f"cannot open: {exc}"[:160]})
            continue
        stem = Path(name).stem
        mapped: list[dict[str, str]] = []
        for ws in src.worksheets:
            title = _safe_sheet_title(f"{ws.title[:16]} ({stem[:12]})", used)
            dst = target.create_sheet(title)
            for row in ws.iter_rows():
                for cell in row:
                    if cell.value is None:
                        continue
                    out = dst.cell(row=cell.row, column=cell.column, value=cell.value)
                    if cell.number_format and cell.number_format != "General":
                        out.number_format = cell.number_format
            mapped.append({"sheet": title, "source_sheet": ws.title})
        sources.append({"file": name, "sheets": mapped, "merged": True})
    if not target.worksheets:
        name, data = files[0]
        return name, data, [{"file": name, "sheets": [], "merged": False}]
    buf = io.BytesIO()
    target.save(buf)
    return "workbooks.xlsx", buf.getvalue(), sources


def _load_excel(task: dict[str, Any]) -> tuple[str, list[dict[str, Any]]]:
    cards = _excel_cards(task)
    if not cards:
        raise FileNotFoundError("excel artifact is missing")
    files = _download_cards(cards)
    if not files:
        raise FileNotFoundError("excel artifact is missing")
    filename, data, sources = _merge_workbooks(files)
    activity = _task_activity(task)
    case_id = str(task.get("case_id") or "")
    session_id = new_session_id()
    session_dir(session_id, create=True)
    suffix = Path(filename).suffix.lower() or ".xlsx"
    relative = f"input{suffix}"
    dest = session_file(session_id, relative)
    dest.write_bytes(data)
    digest = hashlib.sha256(data).hexdigest()
    init_state(
        session_id=session_id,
        file_path=relative,
        file_name=filename,
        file_hash=f"sha256:{digest}",
        file_size=len(data),
        payload={
            "agent_id": "excel_extractor",
            "case_id": case_id,
            "task_id": str(task.get("task_id") or ""),
            "objective": str(task.get("objective") or ""),
            "handoff_message": str(task.get("handoff_message") or ""),
            "activity_base": activity,
            "sources": sources,
            "files": [name for name, _ in files],
        },
    )
    return session_id, sources


# --------------------------------------------------------------------------------------------------
# Inventory for the LLM
# --------------------------------------------------------------------------------------------------


def _sheet_file(sources: list[dict[str, Any]], sheet: str) -> str:
    if len(sources) == 1:
        return str(sources[0].get("file") or "")
    for src in sources:
        for item in src.get("sheets") or []:
            if str(item.get("sheet")) == sheet:
                return str(src.get("file") or "")
    return ""


def _compact_inspect(state: dict[str, Any]) -> dict[str, Any]:
    payload = state.get("payload") if isinstance(state.get("payload"), dict) else {}
    sources = payload.get("sources") if isinstance(payload.get("sources"), list) else []
    intro = execute_tool(state, "workbook_introspect", {})
    meta = intro.get("result") if isinstance(intro, dict) and isinstance(intro.get("result"), dict) else {}
    detect = execute_tool(state, "detect_tables", {"max_tables": 12})
    raw = detect.get("result") if isinstance(detect, dict) else {}
    tables = raw.get("tables") if isinstance(raw, dict) else []
    if not isinstance(tables, list):
        tables = []
    compact: list[dict[str, Any]] = []
    for table in tables[:12]:
        if not isinstance(table, dict) or not table.get("table_id"):
            continue
        cols = table.get("columns") if isinstance(table.get("columns"), list) else []
        sample: list[dict[str, Any]] = []
        described = execute_tool(state, "describe_table", {"table_id": str(table["table_id"]), "sample_rows": 2})
        inner = described.get("result") if isinstance(described, dict) and isinstance(described.get("result"), dict) else {}
        for row in inner.get("sample_rows") or []:
            if isinstance(row, dict):
                sample.append({str(k): _json_safe(v) for k, v in list(row.items())[:24]})
        sheet = str(table.get("sheet") or "")
        compact.append(
            {
                "table_id": table.get("table_id"),
                "file": _sheet_file(sources, sheet),
                "sheet": sheet,
                "range": table.get("range"),
                "row_count": table.get("row_count"),
                "columns": cols[:24],
                "columns_truncated": len(cols) > 24,
                "title": table.get("title") or "",
                "sample": sample,
            }
        )
    sheets = meta.get("sheets") if isinstance(meta.get("sheets"), list) else []
    return {
        "files": [str(s.get("file") or "") for s in sources if s.get("file")],
        "sheets": [
            {
                "name": s.get("name"),
                "rows": s.get("max_row", s.get("rows")),
                "columns": s.get("max_column", s.get("columns")),
                "file": _sheet_file(sources, str(s.get("name") or "")),
            }
            for s in sheets[:40]
            if isinstance(s, dict)
        ],
        "table_count": len(tables),
        "tables": compact,
    }


def _json_safe(value: Any) -> Any:
    if isinstance(value, (datetime, date)):
        return value.strftime("%Y-%m-%d")
    if isinstance(value, float) and value.is_integer():
        return int(value)
    return value


def _parse_jsonish(value: Any) -> Any:
    if isinstance(value, (dict, list)):
        return value
    if isinstance(value, str) and value.strip()[:1] in "{[":
        try:
            return json.loads(value)
        except json.JSONDecodeError:
            return None
    return None


def _engineer_answers(task: dict[str, Any]) -> list[dict[str, Any]]:
    """Compact view of HITL answers for the LLM: question id, chosen option, text — no raw JSON blobs."""
    context = task.get("context") if isinstance(task.get("context"), dict) else {}
    hitl = context.get("hitl") if isinstance(context.get("hitl"), dict) else {}
    answers = hitl.get("answers") if isinstance(hitl.get("answers"), dict) else {}
    out: list[dict[str, Any]] = []
    for key, value in answers.items():
        parsed = _parse_jsonish(value) if not isinstance(value, dict) else value
        row: dict[str, Any] = {"question_id": str(key)}
        if isinstance(parsed, dict):
            for src, dst in (("choice", "choice"), ("label", "label")):
                if str(parsed.get(src) or "").strip():
                    row[dst] = str(parsed[src]).strip()
            text = str(parsed.get("text") or parsed.get("answer") or "").strip()
            if text:
                row["text"] = text[:400]
        elif isinstance(value, str) and value.strip():
            row["text"] = value.strip()[:400]
        else:
            continue
        out.append(row)
    return out[:12]


# --------------------------------------------------------------------------------------------------
# Agent result
# --------------------------------------------------------------------------------------------------


def _agent_result(
    task_id: str,
    status: str,
    message: str,
    *,
    data: dict[str, Any] | None = None,
    artifacts: dict[str, Any] | None = None,
    issues: list[Any] | None = None,
    requests: list[Any] | None = None,
    assumptions: list[Any] | None = None,
) -> dict[str, Any]:
    return {
        "task_id": task_id,
        "agent_id": "excel_extractor",
        "status": status,
        "message": message,
        "data": data or {},
        "artifacts": artifacts or {},
        "issues": issues or [],
        "assumptions": assumptions or [],
        "requests": requests or [],
    }


def _store_result(state: dict[str, Any], result: dict[str, Any]) -> dict[str, Any]:
    state["agent_result"] = result
    save_state(str(state["session_id"]), state)
    return result


def _merge_result(state: dict[str, Any], part: str, summary: str, data: dict[str, Any]) -> dict[str, Any]:
    """Accumulate extraction parts (dates, well parameters) into one completed agent result."""
    payload = state.get("payload") if isinstance(state.get("payload"), dict) else {}
    task_id = str(payload.get("task_id") or "")
    existing = state.get("agent_result") if isinstance(state.get("agent_result"), dict) else {}
    parts = existing.get("data", {}).get("parts") if isinstance(existing.get("data"), dict) else None
    parts = dict(parts) if isinstance(parts, dict) else {}
    parts[part] = summary
    merged: dict[str, Any] = dict(existing.get("data") or {}) if isinstance(existing.get("data"), dict) else {}
    merged.update(data)
    merged["parts"] = parts
    merged["session_id"] = str(state.get("session_id") or "")
    merged["file_name"] = str(state.get("file_name") or "")
    merged["files"] = list(payload.get("files") or [])
    result = _agent_result(
        task_id,
        "completed",
        " ".join(parts[k] for k in ("commissioning", "well_parameters") if k in parts).strip(),
        data=merged,
        artifacts={"excel_session": str(state.get("session_id") or "")},
        assumptions=list(existing.get("assumptions") or []),
    )
    return _store_result(state, result)


def open_session(task: dict[str, Any]) -> dict[str, Any]:
    case_id = str(task.get("case_id") or "")
    task_id = str(task.get("task_id") or "")
    activity = _task_activity(task)
    _emit(case_id, kind="agent.accepted", task_id=task_id, activity=activity, status_message="Разбираю Excel")
    try:
        session_id, sources = _load_excel(task)
    except FileNotFoundError as exc:
        result = _agent_result(
            task_id,
            "needs_input",
            "Нет Excel-файла для извлечения",
            issues=[{"type": "missing_excel", "detail": str(exc)}],
            requests=[
                {
                    "question_id": "Q-excel",
                    "question": "К задаче не приложен Excel-файл с данными. Приложите книгу .xlsx, из которой нужно взять скважины и даты.",
                    "options": [],
                    "accepts": {"free_text": True, "files": ["xlsx"]},
                }
            ],
        )
        _emit(case_id, kind="agent.progress", task_id=task_id, activity=activity, status="needs_input", status_message=result["message"])
        return {"ok": False, "status": "needs_input", "task_id": task_id, "result": result}
    except Exception as exc:
        result = _agent_result(
            task_id,
            "failed",
            str(exc)[:400],
            issues=[{"type": "excel_load_failed", "detail": str(exc)[:400]}],
        )
        _emit(case_id, kind="agent.failed", task_id=task_id, activity=activity, status="failed", status_message=str(exc)[:400])
        return {"ok": False, "status": "failed", "task_id": task_id, "result": result}

    with locked_session(session_id):
        state = load_state(session_id)
        inspect = _compact_inspect(state)
    files = inspect.get("files") or []
    n_sheets = len(inspect.get("sheets") or [])
    n_tables = int(inspect.get("table_count") or 0)
    files_text = "Файл Excel" if len(files) <= 1 else f"Файлов Excel: {len(files)}"
    _emit(
        case_id,
        kind="agent.progress",
        task_id=task_id,
        activity=activity,
        status_message=f"{files_text}: листов {n_sheets}, таблиц {n_tables}. Выбираю таблицу и колонки",
    )
    inputs = task.get("inputs") if isinstance(task.get("inputs"), dict) else {}
    return {
        "ok": True,
        "session_id": session_id,
        "task_id": task_id,
        "file_name": ", ".join(files) if files else str(sources[0].get("file") or "") if sources else "",
        "files": files,
        "inspect": inspect,
        "objective": str(task.get("objective") or ""),
        "handoff_message": str(task.get("handoff_message") or ""),
        "engineer_answers": _engineer_answers(task),
        "rework_reason": str(inputs.get("rework_reason") or ""),
    }


NO_EXTRACT_QUESTION = (
    "Excel Extractor не смог понять, какие данные взять из приложенной книги. "
    "Уточните, на каком листе и в каких колонках находятся скважины и нужные значения (даты ввода, режимы, дебиты)."
)


def session_result(session_id: str) -> dict[str, Any]:
    state = load_state(session_id)
    result = state.get("agent_result")
    if isinstance(result, dict) and result.get("status"):
        return result
    payload = state.get("payload") if isinstance(state.get("payload"), dict) else {}
    return _agent_result(
        str(payload.get("task_id") or ""),
        "needs_input",
        NO_EXTRACT_QUESTION,
        issues=[{"type": "no_extract"}],
        requests=[
            {
                "question_id": "Q-clarify",
                "question": NO_EXTRACT_QUESTION,
                "options": [],
                "accepts": {"free_text": True, "files": ["xlsx"]},
            }
        ],
    )


# --------------------------------------------------------------------------------------------------
# Value normalisation
# --------------------------------------------------------------------------------------------------


def _norm_key(value: Any) -> str:
    return re.sub(r"\s+", " ", str(value or "").strip().lower().replace("ё", "е"))


def _as_well(value: Any) -> str:
    if isinstance(value, bool) or value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    if isinstance(value, int):
        return str(value)
    text = str(value).strip()
    if text.endswith(".0") and text[:-2].isdigit():
        return text[:-2]
    return text


_TNAV_DATE_RE = re.compile(r"^\d{1,2}\s+['\"]?[A-Za-z]{3}['\"]?\s+\d{4}$")


def _as_date(value: Any) -> tuple[Any, bool]:
    """Normalise a cell to a date string; return (value, recognised)."""
    if value in (None, ""):
        return None, False
    if isinstance(value, (datetime, date)):
        return value.strftime("%Y-%m-%d"), True
    if hasattr(value, "strftime") and not isinstance(value, str):
        try:
            return value.strftime("%Y-%m-%d"), True
        except Exception:  # noqa: BLE001
            return str(value), False
    text = str(value).strip()
    if re.fullmatch(r"\d{4}-\d{2}-\d{2}([T ].*)?", text):
        return text[:10], True
    m = re.fullmatch(r"(\d{1,2})[./](\d{1,2})[./](\d{4})", text)
    if m:
        try:
            return date(int(m.group(3)), int(m.group(2)), int(m.group(1))).isoformat(), True
        except ValueError:
            return text, False
    if _TNAV_DATE_RE.match(text):
        return text, True
    return text, False


def _match_column(columns: list[str], wanted: Any) -> str | None:
    key = _norm_key(wanted)
    if not key:
        return None
    for col in columns:
        if _norm_key(col) == key:
            return str(col)
    hits = [str(col) for col in columns if key in _norm_key(col)]
    return hits[0] if len(hits) == 1 else None


def _args_error(code: str, message: str, **details: Any) -> ToolError:
    """Argument-level problems are for the LLM (fix the call), never for the engineer."""
    return ToolError(code, message, details)


def _available_tables(state: dict[str, Any]) -> list[dict[str, Any]]:
    tables = state.get("tables") if isinstance(state.get("tables"), dict) else {}
    return [
        {"table_id": t.get("table_id"), "sheet": t.get("sheet"), "range": t.get("range"), "columns": (t.get("columns") or [])[:24]}
        for t in tables.values()
        if isinstance(t, dict)
    ]


def _resolve_table(state: dict[str, Any], args: dict[str, Any]) -> dict[str, Any]:
    try:
        return _table(state, str(args.get("table_id") or "").strip() or None)
    except ToolError as exc:
        raise _args_error(
            "table_not_found",
            "Такой таблицы в инвентаре нет. Возьми table_id из available_tables (или вызови detect_tables). Это ошибка вызова, инженера спрашивать не нужно.",
            available_tables=_available_tables(state),
        ) from exc


def _repeat_guard(state: dict[str, Any], name: str) -> None:
    history = state.get("tool_history") if isinstance(state.get("tool_history"), list) else []
    calls = sum(1 for item in history if isinstance(item, dict) and item.get("tool") == name)
    if calls >= MAX_TOOL_REPEATS:
        raise _args_error(
            "too_many_attempts",
            f"{name} уже вызывался {calls} раза. Если данных в книге нет — спроси инженера через ask_engineer, иначе заверши ответ.",
        )


# --------------------------------------------------------------------------------------------------
# Extraction tools (LLM chooses table and columns; extraction itself is deterministic)
# --------------------------------------------------------------------------------------------------


def _wells_preview(wells: list[str], limit: int = 6) -> str:
    if len(wells) <= limit:
        return ", ".join(wells)
    return ", ".join(wells[:limit]) + f" и ещё {len(wells) - limit}"


def _plural(n: int, one: str, few: str, many: str) -> str:
    """Russian plural: 1 скважина, 2 скважины, 5 скважин."""
    n10, n100 = n % 10, n % 100
    if n10 == 1 and n100 != 11:
        return f"{n} {one}"
    if 2 <= n10 <= 4 and not 12 <= n100 <= 14:
        return f"{n} {few}"
    return f"{n} {many}"


_FIELD_LABELS = {
    "date": "дата ввода",
    "group": "группа",
    "phase": "фаза",
    "i": "ячейка I",
    "j": "ячейка J",
    "md_top": "интервал MD",
    "md_bot": "интервал MD",
    "diameter": "диаметр",
    "control": "режим",
    "rate": "дебит",
    "bhp": "лимит BHP",
    "thp": "лимит THP",
    "vfp_table": "таблица VFP",
    "welltrack_include": "файл траектории",
}


@tool(
    _schema(
        "extract_commissioning",
        "Extract well → commissioning date facts from one table the model selected in the inventory. Deterministic; validates that the chosen columns exist and hold dates.",
        {
            "table_id": {"type": "string", "minLength": 1},
            "well_column": {"type": "string", "minLength": 1},
            "date_column": {"type": "string", "minLength": 1},
        },
        ["table_id", "well_column", "date_column"],
    )
)
def extract_commissioning_tool(ctx: dict[str, Any], args: dict[str, Any]) -> dict[str, Any]:
    state = ctx["state"]
    _repeat_guard(state, "extract_commissioning")
    missing = [k for k in ("table_id", "well_column", "date_column") if not str(args.get(k) or "").strip()]
    if missing:
        raise _args_error(
            "spec_incomplete",
            "Укажи table_id, well_column и date_column из инвентаря (inspect.tables: columns и sample). Это ошибка вызова, инженера спрашивать не нужно.",
            missing=missing,
            where_to_find={
                "table_id": "inspect.tables[].table_id",
                "well_column": "колонка со скважинами в inspect.tables[].columns",
                "date_column": "колонка с новой датой ввода (не baseline) в inspect.tables[].columns",
            },
            available_tables=_available_tables(state),
        )
    table = _resolve_table(state, args)
    columns = [str(c) for c in (table.get("columns") or [])]
    well_col = _match_column(columns, args.get("well_column"))
    date_col = _match_column(columns, args.get("date_column"))
    bad = {name: str(args.get(name)) for name, col in (("well_column", well_col), ("date_column", date_col)) if not col}
    if bad:
        raise _args_error(
            "column_not_found",
            "Колонки с такими именами в таблице нет. Возьми точное имя из available_columns. Это ошибка вызова, инженера спрашивать не нужно.",
            not_found=bad,
            available_columns=columns,
        )
    if well_col == date_col:
        raise _args_error("column_not_found", "well_column и date_column должны быть разными колонками.", available_columns=columns)
    facts: list[dict[str, Any]] = []
    seen: set[str] = set()
    non_empty = 0
    unrecognised: list[str] = []
    skipped_empty = 0
    for row in _rows_for_table(state, table):
        rec = _record(columns, row)
        well = _as_well(rec.get(well_col))
        raw_date = rec.get(date_col)
        if raw_date in (None, ""):
            skipped_empty += 1
            continue
        value, recognised = _as_date(raw_date)
        non_empty += 1
        if not recognised:
            if len(unrecognised) < 5:
                unrecognised.append(str(raw_date)[:40])
            continue
        if not well:
            skipped_empty += 1
            continue
        key = well.casefold()
        if key in seen:
            continue
        seen.add(key)
        facts.append({"well": well, "date": value})
    if non_empty and len(unrecognised) * 2 > non_empty:
        raise _args_error(
            "column_not_dates",
            "Выбранная колонка не содержит дат. Выбери другую колонку (см. sample) или спроси инженера, если колонки с датами нет.",
            date_column=date_col,
            sample_values=unrecognised,
            available_columns=columns,
        )
    if not facts:
        raise _args_error(
            "no_rows",
            "В выбранных колонках нет ни одной пары «скважина — дата». Проверь колонки по sample или выбери другую таблицу.",
            available_columns=columns,
        )
    dates = sorted(str(f["date"]) for f in facts)
    wells = [str(f["well"]) for f in facts]
    span = f"на {dates[0]}" if dates[0] == dates[-1] else f"с {dates[0]} по {dates[-1]}"
    summary = f"Даты ввода: {_plural(len(facts), 'скважина', 'скважины', 'скважин')} ({_wells_preview(wells)}), {span}."
    data = {
        "facts": facts,
        "preview": facts[:10],
        "total_count": len(facts),
        "excel_table": table.get("table_id"),
        "table_ids": [table.get("table_id")],
        "commissioning_source": {
            "table_id": table.get("table_id"),
            "sheet": table.get("sheet"),
            "range": table.get("range"),
            "well_column": well_col,
            "date_column": date_col,
            "rows_skipped": skipped_empty,
            "unrecognised_dates": unrecognised,
        },
    }
    result = _merge_result(state, "commissioning", summary, data)
    _emit_from_state(state, summary, status="completed")
    return {"status": "completed", "message": summary, "facts_count": len(facts), "preview": facts[:5], "result": {"status": result["status"]}}


def _mapping_from_args(raw: Any) -> dict[str, str]:
    parsed = _parse_jsonish(raw) if not isinstance(raw, dict) else raw
    if not isinstance(parsed, dict):
        return {}
    out: dict[str, str] = {}
    for key, value in parsed.items():
        field = _norm_key(key).replace(" ", "_")
        if field in NEW_WELL_FIELDS and str(value or "").strip():
            out[field] = str(value).strip()
    return out


def _cell_value(value: Any) -> Any:
    if isinstance(value, (datetime, date)):
        return value.strftime("%Y-%m-%d")
    if isinstance(value, float) and value.is_integer():
        return int(value)
    if isinstance(value, str):
        return value.strip()
    return value


@tool(
    _schema(
        "extract_well_parameters",
        "Extract a per-well parameters table (group, MD interval, diameter, control, rate, BHP, VFP, trajectory file…) for new wells. Model selects the table, the well column and an optional column→field mapping; extraction is deterministic.",
        {
            "table_id": {"type": "string", "minLength": 1},
            "well_column": {"type": "string", "minLength": 1},
            "mapping": {"type": ["object", "string"]},
        },
        ["table_id", "well_column"],
    )
)
def extract_well_parameters_tool(ctx: dict[str, Any], args: dict[str, Any]) -> dict[str, Any]:
    state = ctx["state"]
    _repeat_guard(state, "extract_well_parameters")
    missing = [k for k in ("table_id", "well_column") if not str(args.get(k) or "").strip()]
    if missing:
        raise _args_error(
            "spec_incomplete",
            "Укажи table_id и well_column из инвентаря; mapping — JSON {поле: колонка} для полей "
            + ", ".join(NEW_WELL_FIELDS[1:])
            + ". Это ошибка вызова, инженера спрашивать не нужно.",
            missing=missing,
            available_tables=_available_tables(state),
        )
    table = _resolve_table(state, args)
    columns = [str(c) for c in (table.get("columns") or [])]
    well_col = _match_column(columns, args.get("well_column"))
    if not well_col:
        raise _args_error(
            "column_not_found",
            "Колонки со скважинами с таким именем нет. Возьми точное имя из available_columns.",
            not_found={"well_column": str(args.get("well_column"))},
            available_columns=columns,
        )
    mapping = _mapping_from_args(args.get("mapping"))
    resolved_map: dict[str, str] = {}
    bad: dict[str, str] = {}
    for field, col_name in mapping.items():
        col = _match_column(columns, col_name)
        if col:
            resolved_map[field] = col
        else:
            bad[field] = col_name
    if bad:
        raise _args_error(
            "column_not_found",
            "В mapping указаны колонки, которых нет в таблице. Возьми точные имена из available_columns.",
            not_found=bad,
            available_columns=columns,
        )
    mapped_cols = set(resolved_map.values()) | {well_col}
    rows_out: list[dict[str, Any]] = []
    seen: set[str] = set()
    for row in _rows_for_table(state, table):
        rec = _record(columns, row)
        well = _as_well(rec.get(well_col))
        if not well or well.casefold() in seen:
            continue
        seen.add(well.casefold())
        out: dict[str, Any] = {"well": well}
        for field, col in resolved_map.items():
            value = rec.get(col)
            if value in (None, ""):
                continue
            out[field] = _as_date(value)[0] if field == "date" else _cell_value(value)
        for col in columns:
            if col in mapped_cols:
                continue
            value = rec.get(col)
            if value in (None, ""):
                continue
            out[col] = _cell_value(value)
        rows_out.append(out)
    if not rows_out:
        raise _args_error("no_rows", "В таблице нет строк со скважинами в выбранной колонке.", available_columns=columns)
    wells = [str(r["well"]) for r in rows_out]
    labels: list[str] = []
    for field in resolved_map:
        label = _FIELD_LABELS.get(field, "")
        if label and label not in labels:
            labels.append(label)
    extra = len([c for c in columns if c not in mapped_cols])
    described = ", ".join(labels) if labels else ""
    if extra:
        described = (described + ", " if described else "") + _plural(extra, "прочая колонка", "прочие колонки", "прочих колонок")
    summary = f"Параметры новых скважин: {_plural(len(rows_out), 'скважина', 'скважины', 'скважин')} ({_wells_preview(wells)})" + (
        f" — {described}." if described else "."
    )
    data = {
        "new_wells": rows_out,
        "new_wells_source": {
            "table_id": table.get("table_id"),
            "sheet": table.get("sheet"),
            "range": table.get("range"),
            "well_column": well_col,
            "mapping": resolved_map,
            "columns": columns,
        },
    }
    result = _merge_result(state, "well_parameters", summary, data)
    _emit_from_state(state, summary, status="completed")
    return {"status": "completed", "message": summary, "wells_count": len(rows_out), "preview": rows_out[:3], "result": {"status": result["status"]}}


# --------------------------------------------------------------------------------------------------
# ask_engineer — the only way the LLM asks a human (prose, validated)
# --------------------------------------------------------------------------------------------------

_MACHINE_TOKEN_RE = re.compile(r"[a-z]+_[a-z_]+|[a-z_]+=[a-z0-9]+|[{}\[\]<>]|\w\|\w")


def human_text_problems(text: str) -> list[str]:
    """Why a string is not fit for an engineer: snake_case ids, key=value, JSON braces, a|b enums."""
    problems: list[str] = []
    stripped = str(text or "").strip()
    if len(stripped) < 12:
        problems.append("слишком коротко для вопроса инженеру")
    if not re.search(r"[А-Яа-яЁё]{3,}", stripped):
        problems.append("вопрос должен быть по-русски")
    tokens = sorted({m.group(0) for m in _MACHINE_TOKEN_RE.finditer(stripped)})
    if tokens:
        problems.append("машинные токены: " + ", ".join(tokens[:6]))
    return problems


def _options_for_human(raw: Any) -> list[dict[str, str]]:
    items = raw
    if isinstance(raw, str):
        parsed = _parse_jsonish(raw)
        items = parsed if isinstance(parsed, list) else [part.strip() for part in raw.split(";") if part.strip()]
    if not isinstance(items, list):
        return []
    out: list[dict[str, str]] = []
    for item in items:
        if isinstance(item, dict):
            label = str(item.get("label") or item.get("text") or item.get("value") or "").strip()
            value = str(item.get("value") or label).strip()
        else:
            label = str(item).strip()
            value = label
        if label:
            out.append({"value": value, "label": label})
    return out[:8]


@tool(
    _schema(
        "ask_engineer",
        "Ask the engineer one question in plain Russian when the inventory cannot resolve which table/column to use. Options are shown as buttons.",
        {
            "question": {"type": "string", "minLength": 1},
            "options": {"type": ["array", "string"]},
            "topic": {"type": "string"},
        },
        ["question"],
    )
)
def ask_engineer_tool(ctx: dict[str, Any], args: dict[str, Any]) -> dict[str, Any]:
    state = ctx["state"]
    question = str(args.get("question") or "").strip()
    problems = human_text_problems(question)
    options = _options_for_human(args.get("options"))
    for opt in options:
        problems.extend(f"вариант «{opt['label']}»: {p}" for p in human_text_problems(opt["label"]) if "коротко" not in p and "по-русски" not in p)
    if problems:
        raise _args_error(
            "question_not_human",
            "Переформулируй вопрос для инженера: обычная русская фраза, что именно нужно и зачем; "
            "варианты — как их видит инженер (названия листов и колонок), без имён полей и JSON.",
            problems=problems,
        )
    payload = state.get("payload") if isinstance(state.get("payload"), dict) else {}
    topic = re.sub(r"[^a-z0-9_]+", "_", str(args.get("topic") or "excel").strip().lower()).strip("_") or "excel"
    existing = state.get("agent_result") if isinstance(state.get("agent_result"), dict) else {}
    result = _agent_result(
        str(payload.get("task_id") or ""),
        "needs_input",
        question,
        data={**(existing.get("data") if isinstance(existing.get("data"), dict) else {}), "asked_by": "excel_extractor_llm"},
        artifacts={"excel_session": str(state.get("session_id") or "")},
        issues=[{"type": "engineer_input_required", "topic": topic}],
        requests=[
            {
                "question_id": f"Q-{topic}",
                "question": question,
                "required": True,
                "type": "choice" if options else "text",
                "options": options,
                "accepts": {"free_text": True, "files": ["xlsx"]},
            }
        ],
    )
    _store_result(state, result)
    _emit_from_state(state, "Нужно уточнение инженера: " + question[:200], status="needs_input")
    return {"status": "needs_input", "message": question, "result": {"status": "needs_input"}}
