"""Excel Extractor as a MAS agent — the ``AgentService`` the n8n workflow «Agent — Excel Extractor» talks to.

* ``open_session(agent_task)`` loads **every** Excel attachment of the case into one session (several
  ``.xlsx`` merged into one workbook) and returns a compact **inventory** for the LLM: files, sheets,
  tables with columns and two sample rows. No decision is taken here — no keyword heuristics pick the
  table or the column with wells or dates.
* ``result(state)`` returns what the extraction tools fixed, or a ``needs_input`` in prose when the LLM
  finished without extracting anything.
* ``normalize_args`` accepts a few unambiguous transport aliases of the n8n HTTP Tool; ``after_tool``
  writes progress lines to the Activity feed.

LLM tools: ``excel_tools.py`` (workbook discovery), ``state_tools.py`` (session state),
``agent_tools.py`` (``extract_commissioning``, ``extract_well_parameters``, ``ask_engineer``).
Same shape as ``agents-template/demo_agent/app/agent.py``. No SCHEDULE text is ever written here.
"""

from __future__ import annotations

import hashlib
import logging
import os
import re
from datetime import date, datetime
from pathlib import Path
from typing import Any

from mas_agent_kit import ActivityClient, AgentService, CasePacket

from .sessions import STORE, init_state, session_dir, session_file
from .tools import TOOLS, execute_tool

logger = logging.getLogger(__name__)
AGENT_ID = "excel_extractor"
EXCEL_SUFFIXES = (".xlsx", ".xlsm", ".xltx", ".xltm", ".xls")
MERGEABLE_SUFFIXES = (".xlsx", ".xlsm", ".xltx", ".xltm")

MISSING_EXCEL_QUESTION = "К задаче не приложен Excel-файл с данными. Приложите книгу .xlsx, из которой нужно взять скважины и даты."
NO_EXTRACT_QUESTION = (
    "Excel Extractor не смог понять, какие данные взять из приложенной книги. "
    "Уточните, на каком листе и в каких колонках находятся скважины и нужные значения (даты ввода, режимы, дебиты)."
)
# Feed lines while the LLM explores the workbook; extraction tools post their own precise lines.
TOOL_PROGRESS = {
    "match_tables": "Подбираю таблицу под запрос",
    "describe_table": "Смотрю колонки таблицы",
    "list_column_values": "Читаю значения колонки",
    "sheet_preview": "Смотрю фрагмент листа",
}
TOOL_ERROR_PROGRESS = {
    "column_not_found": "Уточняю колонки таблицы",
    "table_not_found": "Уточняю таблицу",
    "column_not_dates": "Выбранная колонка не похожа на даты — подбираю другую",
    "spec_incomplete": "Уточняю, какую таблицу и колонки взять",
    "question_not_human": "Формулирую вопрос инженеру",
}


def excel_request(question_id: str, question: str) -> dict[str, Any]:
    return {"question_id": question_id, "question": question, "options": [], "accepts": {"free_text": True, "files": ["xlsx"]}}


class ExcelExtractorAgent(AgentService):
    agent_id = AGENT_ID
    store = STORE
    tools = TOOLS
    activity_base_url = os.getenv("ACTIVITY_BASE_URL", "").rstrip("/")

    # -- session --------------------------------------------------------------------------------

    def open_session(self, task: dict[str, Any]) -> dict[str, Any]:
        activity = self.activity(task)
        activity.accepted("Разбираю Excel")
        stub = {"task_id": str(task.get("task_id") or "")}
        try:
            state, sources = self._load_excel(task, activity)
        except FileNotFoundError as exc:
            result = self.new_result(stub, "needs_input", "Нет Excel-файла для извлечения", issues=[{"type": "missing_excel", "detail": str(exc)}], requests=[excel_request("Q-excel", MISSING_EXCEL_QUESTION)])
            activity.progress(result["message"], status="needs_input")
            return self.not_opened(result)
        except Exception as exc:  # noqa: BLE001 — a broken workbook is a final answer, not a crash
            activity.failed(str(exc)[:400])
            return self.not_opened(self.new_result(stub, "failed", str(exc)[:400], issues=[{"type": "excel_load_failed", "detail": str(exc)[:400]}]))

        with self.store.lock(state["session_id"]):
            state = self.store.load(state["session_id"])
            inspect = compact_inspect(state)
        files = inspect.get("files") or []
        files_text = "Файл Excel" if len(files) <= 1 else f"Файлов Excel: {len(files)}"
        activity.progress(f"{files_text}: листов {len(inspect.get('sheets') or [])}, таблиц {int(inspect.get('table_count') or 0)}. Выбираю таблицу и колонки")
        return self.opened(
            state,
            file_name=", ".join(files) if files else str(sources[0].get("file") or "") if sources else "",
            files=files,
            inspect=inspect,
        )

    def result(self, state: dict[str, Any]) -> dict[str, Any]:
        result = state.get("result")
        if isinstance(result, dict) and result.get("status"):
            return result
        return self.new_result(state, "needs_input", NO_EXTRACT_QUESTION, issues=[{"type": "no_extract"}], requests=[excel_request("Q-clarify", NO_EXTRACT_QUESTION)])

    # -- hooks ----------------------------------------------------------------------------------

    def normalize_args(self, tool_name: str, args: dict[str, Any]) -> dict[str, Any]:
        """Accept a few unambiguous LLM transport aliases before strict tool validation.

        Protects the named agent-tool transport from common function-calling vocabulary such as
        ``fields`` versus the API's canonical ``select``. Never invents an id, a filter value or a column.
        """
        normalized = dict(args)

        if tool_name == "describe_table" and "table_id" not in normalized:
            candidate = normalized.get("verified_table")
            if isinstance(candidate, str):
                normalized["table_id"] = candidate
            elif isinstance(candidate, dict) and isinstance(candidate.get("table_id"), str):
                normalized["table_id"] = candidate["table_id"]

        if tool_name in {"query_table", "detect_tables", "describe_table", "list_column_values"}:
            for field in ("limit", "max_tables", "sample_rows"):
                raw = normalized.get(field)
                if isinstance(raw, str) and raw.strip().isdigit():
                    normalized[field] = int(raw.strip())
                elif isinstance(raw, bool):
                    continue
                elif isinstance(raw, float) and raw.is_integer():
                    normalized[field] = int(raw)

        if tool_name == "query_table":
            if "table_id" not in normalized and isinstance(normalized.get("table"), str):
                normalized["table_id"] = normalized["table"]
            if "select" not in normalized and isinstance(normalized.get("fields"), list):
                normalized["select"] = normalized["fields"]
            # Preserve omission of optional fields: injecting ``None`` would turn a valid no-filter
            # query into ``filters: null`` for the strict validator.
            for field in ("select", "filters"):
                if field in normalized:
                    normalized[field] = _n8n_json_sequence(normalized[field])
            filters = normalized.get("filters")
            if isinstance(filters, dict):
                normalized["filters"] = _canonical_filters(filters)

        if tool_name == "save_agent_plan":
            candidate = normalized.get("verified_table")
            table_id = candidate.get("table_id") if isinstance(candidate, dict) else candidate
            if "selected_table_ids" not in normalized and isinstance(table_id, str):
                normalized["selected_table_ids"] = [table_id]
            if "field_mapping" not in normalized and isinstance(normalized.get("mapping"), dict):
                normalized["field_mapping"] = normalized["mapping"]
            if "select" not in normalized and isinstance(normalized.get("fields"), list):
                normalized["select"] = normalized["fields"]
            if "plan" not in normalized and isinstance(table_id, str):
                normalized["plan"] = f"Selected verified table {table_id}"
            for field in ("selected_table_ids", "select", "filters", "assumptions", "warnings"):
                if field in normalized:
                    normalized[field] = _n8n_json_sequence(normalized[field])
            filters = normalized.get("filters")
            if isinstance(filters, dict):
                normalized["filters"] = _canonical_filters(filters)

        return normalized

    def after_tool(self, state: dict[str, Any], tool_name: str, result: dict[str, Any]) -> None:
        """Live Activity lines while the n8n AI Agent calls the discovery tools."""
        message = TOOL_PROGRESS.get(tool_name, "")
        if tool_name == "workbook_introspect":
            sheets = result.get("sheets") if isinstance(result.get("sheets"), list) else []
            message = f"Смотрю структуру Excel: листов {len(sheets)}"
        elif tool_name == "detect_tables":
            tables = result.get("tables") or result.get("detected_tables") or []
            message = f"Нашёл таблиц: {len(tables) if isinstance(tables, list) else 0}"
        elif tool_name == "query_table":
            rows = result.get("preview") or result.get("rows") or result.get("records") or []
            n = result.get("row_count") if isinstance(result.get("row_count"), int) else (len(rows) if isinstance(rows, list) else 0)
            message = f"Читаю таблицу: {n} строк"
        elif tool_name in {"extract_commissioning", "extract_well_parameters", "ask_engineer"} and not result.get("ok", True):
            message = TOOL_ERROR_PROGRESS.get(str(result.get("code") or ""), "")
        if message:
            self.activity_for_state(state).progress(message)

    # -- workbooks of the case → one session ----------------------------------------------------

    def excel_cards(self, task: dict[str, Any]) -> list[dict[str, Any]]:
        """Every Excel workbook attached to the case: explicit path, explicit URL, artifact cards, case state."""
        inputs = task.get("inputs") if isinstance(task.get("inputs"), dict) else {}
        cards: list[dict[str, Any]] = []
        path = inputs.get("excel_path")
        if path and Path(str(path)).is_file():
            cards.append({"artifact_id": "excel", "filename": Path(str(path)).name, "path": str(path)})
        url = inputs.get("excel_url") or inputs.get("artifact_url")
        if url:
            cards.append({"artifact_id": "excel", "filename": str(inputs.get("excel_filename") or ""), "url": str(url)})
        if cards:
            return cards
        activity = self.activity(task)
        packet = CasePacket.load(task, activity)
        found: dict[str, dict[str, Any]] = {}
        for aid, card in packet.artifacts.items():
            if _is_excel_card(card):
                found[aid] = {"artifact_id": aid, "filename": str(card.get("filename") or "")}
        ids = inputs.get("artifact_ids") if isinstance(inputs.get("artifact_ids"), list) else []
        for aid in ids:
            aid = str(aid or "").strip()
            if re.fullmatch(r"excel(_\d+)?", aid) and aid not in found:
                found[aid] = {"artifact_id": aid, "filename": ""}
        if activity.configured and "excel" not in found:
            # The first workbook of a case always has the fixed id `excel`; older orchestrator states may
            # list only later workbooks (or none at all). Probe it — a missing artifact is skipped on download.
            found["excel"] = {"artifact_id": "excel", "filename": ""}

        def order(card: dict[str, Any]) -> tuple[int, str]:
            m = re.fullmatch(r"excel_(\d+)", card["artifact_id"])
            return (int(m.group(1)) if m else 0, card["artifact_id"])

        for card in sorted(found.values(), key=order):
            if activity.configured:
                card["url"] = activity.artifact_url(card["artifact_id"])
            cards.append(card)
        return cards

    def _load_excel(self, task: dict[str, Any], activity: ActivityClient) -> tuple[dict[str, Any], list[dict[str, Any]]]:
        cards = self.excel_cards(task)
        if not cards:
            raise FileNotFoundError("excel artifact is missing")
        files = _download_cards(cards, activity)
        if not files:
            raise FileNotFoundError("excel artifact is missing")
        filename, data, sources = _merge_workbooks(files)
        session_id = self.store.new_id()
        session_dir(session_id, create=True)
        relative = "input" + (Path(filename).suffix.lower() or ".xlsx")
        session_file(session_id, relative).write_bytes(data)
        state = init_state(
            session_id=session_id,
            file_path=relative,
            file_name=filename,
            file_hash=f"sha256:{hashlib.sha256(data).hexdigest()}",
            file_size=len(data),
            payload={"agent_id": self.agent_id, "sources": sources, "files": [name for name, _ in files]},
            # Case fields the kit helpers read (``new_result``, ``activity_for_state``).
            agent_id=self.agent_id,
            case_id=str(task.get("case_id") or ""),
            task_id=str(task.get("task_id") or ""),
            objective=str(task.get("objective") or ""),
            handoff_message=str(task.get("handoff_message") or ""),
            activity_base_url=activity.base_url,
            inputs=task.get("inputs") if isinstance(task.get("inputs"), dict) else {},
            context=task.get("context") if isinstance(task.get("context"), dict) else {},
            result=None,
        )
        return state, sources


# --------------------------------------------------------------------------------------------------
# Helpers: workbook cards, merging, inventory, n8n transport quirks
# --------------------------------------------------------------------------------------------------


def _is_excel_name(name: Any) -> bool:
    return str(name or "").lower().endswith(EXCEL_SUFFIXES)


def _is_excel_card(card: dict[str, Any]) -> bool:
    aid = str(card.get("artifact_id") or "")
    role = str(card.get("role") or "")
    return role == "excel" or (not role and (bool(re.fullmatch(r"excel(_\d+)?", aid)) or _is_excel_name(card.get("filename"))))


def _download_cards(cards: list[dict[str, Any]], activity: ActivityClient) -> list[tuple[str, bytes]]:
    files: list[tuple[str, bytes]] = []
    errors: list[str] = []
    for card in cards:
        data: bytes | None = None
        name = str(card.get("filename") or "")
        if card.get("path"):
            data = Path(str(card["path"])).read_bytes()
            name = name or Path(str(card["path"])).name
        elif card.get("url"):
            try:
                data, header_name = activity.download_url(str(card["url"]), timeout=30)
            except Exception as exc:  # noqa: BLE001
                logger.warning("excel artifact %s skipped: %s", card.get("artifact_id"), exc)
                errors.append(f"{card.get('artifact_id')}: {exc}")
                continue
            name = name or header_name
        if not data:
            continue
        if not _is_excel_name(name):
            name = (name or str(card.get("artifact_id") or "workbook")) + ".xlsx"
        files.append((name, data))
    if not files and errors:
        raise FileNotFoundError("excel artifact fetch failed: " + "; ".join(errors)[:400])
    return files


def _safe_sheet_title(title: str, used: set[str]) -> str:
    clean = re.sub(r"[\[\]:*?/\\]", " ", title).strip() or "Sheet"
    base = clean[:31]
    candidate = base
    n = 2
    while candidate in used:
        suffix = f" ({n})"
        candidate = base[: 31 - len(suffix)] + suffix
        n += 1
    used.add(candidate)
    return candidate


def _merge_workbooks(files: list[tuple[str, bytes]]) -> tuple[str, bytes, list[dict[str, Any]]]:
    """One session file for all attachments.

    A single workbook is stored as-is. Several .xlsx are merged into one workbook (values only) with
    sheet titles ``<sheet> (<file stem>)`` so the inventory tells the LLM which file a table came from.
    """
    if len(files) == 1:
        name, data = files[0]
        return name, data, [{"file": name, "sheets": [], "merged": False}]
    import io

    from openpyxl import Workbook, load_workbook

    target = Workbook()
    target.remove(target.active)
    used: set[str] = set()
    sources: list[dict[str, Any]] = []
    for name, data in files:
        if not name.lower().endswith(MERGEABLE_SUFFIXES):
            sources.append({"file": name, "sheets": [], "merged": False, "skipped": "only .xlsx workbooks are merged"})
            continue
        try:
            src = load_workbook(io.BytesIO(data), data_only=True)
        except Exception as exc:  # noqa: BLE001
            sources.append({"file": name, "sheets": [], "merged": False, "skipped": f"cannot open: {exc}"[:160]})
            continue
        stem = Path(name).stem
        mapped: list[dict[str, str]] = []
        for ws in src.worksheets:
            title = _safe_sheet_title(f"{ws.title[:16]} ({stem[:12]})", used)
            dst = target.create_sheet(title)
            for row in ws.iter_rows():
                for cell in row:
                    if cell.value is None:
                        continue
                    out = dst.cell(row=cell.row, column=cell.column, value=cell.value)
                    if cell.number_format and cell.number_format != "General":
                        out.number_format = cell.number_format
            mapped.append({"sheet": title, "source_sheet": ws.title})
        sources.append({"file": name, "sheets": mapped, "merged": True})
    if not target.worksheets:
        name, data = files[0]
        return name, data, [{"file": name, "sheets": [], "merged": False}]
    buf = io.BytesIO()
    target.save(buf)
    return "workbooks.xlsx", buf.getvalue(), sources


def _sheet_file(sources: list[dict[str, Any]], sheet: str) -> str:
    if len(sources) == 1:
        return str(sources[0].get("file") or "")
    for src in sources:
        for item in src.get("sheets") or []:
            if str(item.get("sheet")) == sheet:
                return str(src.get("file") or "")
    return ""


def _json_safe(value: Any) -> Any:
    if isinstance(value, (datetime, date)):
        return value.strftime("%Y-%m-%d")
    if isinstance(value, float) and value.is_integer():
        return int(value)
    return value


def compact_inspect(state: dict[str, Any]) -> dict[str, Any]:
    """Inventory for the LLM: files, sheets, up to 12 tables with columns and two sample rows."""
    payload = state.get("payload") if isinstance(state.get("payload"), dict) else {}
    sources = payload.get("sources") if isinstance(payload.get("sources"), list) else []
    meta = execute_tool(state, "workbook_introspect", {})
    tables = execute_tool(state, "detect_tables", {"max_tables": 12}).get("tables")
    if not isinstance(tables, list):
        tables = []
    compact: list[dict[str, Any]] = []
    for table in tables[:12]:
        if not isinstance(table, dict) or not table.get("table_id"):
            continue
        cols = table.get("columns") if isinstance(table.get("columns"), list) else []
        sample: list[dict[str, Any]] = []
        described = execute_tool(state, "describe_table", {"table_id": str(table["table_id"]), "sample_rows": 2})
        for row in described.get("sample_rows") or []:
            if isinstance(row, dict):
                sample.append({str(k): _json_safe(v) for k, v in list(row.items())[:24]})
        sheet = str(table.get("sheet") or "")
        compact.append(
            {
                "table_id": table.get("table_id"),
                "file": _sheet_file(sources, sheet),
                "sheet": sheet,
                "range": table.get("range"),
                "row_count": table.get("row_count"),
                "columns": cols[:24],
                "columns_truncated": len(cols) > 24,
                "title": table.get("title") or "",
                "sample": sample,
            }
        )
    sheets = meta.get("sheets") if isinstance(meta.get("sheets"), list) else []
    return {
        "files": [str(s.get("file") or "") for s in sources if s.get("file")],
        "sheets": [
            {"name": s.get("name"), "rows": s.get("max_row", s.get("rows")), "columns": s.get("max_column", s.get("columns")), "file": _sheet_file(sources, str(s.get("name") or ""))}
            for s in sheets[:40]
            if isinstance(s, dict)
        ],
        "table_count": len(tables),
        "tables": compact,
    }


def _n8n_json_sequence(value: Any) -> Any:
    """Turn n8n HTTP Tool 1.1's object-shaped JSON (``{"0": …, "1": …}``) into a canonical array.

    The node's UI exposes JSON parameters, but its generated Zod schema accepts JSON objects rather than
    arrays. Kept entirely at the n8n adapter boundary; regular clients continue to use arrays.
    """
    if not isinstance(value, dict):
        return value
    if not value:
        return []
    keys = list(value)
    if not all(isinstance(key, str) and key.isdigit() for key in keys):
        return value
    return [value[str(index)] for index in sorted(int(key) for key in keys)]


def _canonical_filters(filters: dict[str, Any]) -> list[dict[str, Any]]:
    """``{"well": "101"}`` / ``{"well": {"operator": "eq", "value": "101"}}`` → ``[{"field", "operator", "value"}]``."""
    return [
        {"field": field, **condition} if isinstance(condition, dict) else {"field": field, "operator": "eq", "value": condition}
        for field, condition in filters.items()
        if isinstance(field, str)
    ]


agent = ExcelExtractorAgent()

# Register the LLM tools on ``agent.tools`` (these modules never invoke an LLM).
from . import agent_tools, excel_tools, state_tools  # noqa: E402,F401
